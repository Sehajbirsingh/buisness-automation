import os
import pickle
import base64
from datetime import datetime, timedelta
from email.mime.text import MIMEText

from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build

SCOPES = [
    'https://www.googleapis.com/auth/gmail.readonly',
    'https://www.googleapis.com/auth/gmail.send',
    'https://www.googleapis.com/auth/calendar'
]

def authenticate_google():
    creds = None
    if os.path.exists('token.pkl'):
        with open('token.pkl', 'rb') as token:
            creds = pickle.load(token)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(
                'client_secret.json', SCOPES, redirect_uri='urn:ietf:wg:oauth:2.0:oob'
            )
            auth_url, _ = flow.authorization_url(prompt='consent')
            print("🔐 Authorize your app:\n", auth_url)
            code = input("Paste code: ")
            flow.fetch_token(code=code)
            creds = flow.credentials
        with open('token.pkl', 'wb') as token:
            pickle.dump(creds, token)
    return creds

# Build services
creds = authenticate_google()
gmail_service = build('gmail', 'v1', credentials=creds)
calendar_service = build('calendar', 'v3', credentials=creds)

from langchain_google_genai import ChatGoogleGenerativeAI, GoogleGenerativeAIEmbeddings

# Load Gemini Model
os.environ["GOOGLE_API_KEY"] = "AIzaSyD4a5CQZDdPj9wGlBXJy4pnOCkbBCXPZcI"
llm = ChatGoogleGenerativeAI(model="gemini-1.5-pro", temperature=0.2)
embedding = GoogleGenerativeAIEmbeddings(model="models/embedding-001")

from langchain.document_loaders import PyPDFLoader, UnstructuredWordDocumentLoader, CSVLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.vectorstores import FAISS
from langchain.chains import RetrievalQA

def load_documents(files):
    docs = []
    for file in files:
        if file.endswith(".pdf"):
            docs.extend(PyPDFLoader(file).load())
        elif file.endswith(".docx"):
            docs.extend(UnstructuredWordDocumentLoader(file).load())
        elif file.endswith(".csv"):
            docs.extend(CSVLoader(file_path=file).load())
    return docs

def split_documents(documents):
    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
    return splitter.split_documents(documents)

def build_vectorstore(docs):
    texts = split_documents(docs)
    return FAISS.from_documents(texts, embedding)

def generate_report_with_context(vectorstore, prompt):
    retriever = vectorstore.as_retriever(search_type="similarity", k=4)
    qa = RetrievalQA.from_chain_type(llm=llm, retriever=retriever)
    return qa.run(prompt)

from docx import Document
from fpdf import FPDF

def export_to_docx(text, filename="report.docx"):
    doc = Document()
    doc.add_heading("AI-Generated Report", level=1)
    doc.add_paragraph(text)
    doc.save(filename)

def export_to_pdf(text, filename="report.pdf"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in text.split("\n"):
        pdf.multi_cell(0, 10, txt=line)
    pdf.output(filename)

def send_email(recipient, subject, body):
    message = MIMEText(body, _charset='UTF-8')
    message['to'] = recipient
    message['from'] = "me"
    message['subject'] = subject
    raw = base64.urlsafe_b64encode(message.as_bytes()).decode()
    try:
        gmail_service.users().messages().send(
            userId='me', body={'raw': raw}
        ).execute()
    except Exception as e:
        print("❌ Email send failed:", e)

def schedule_meeting(title, datetime_str, duration_minutes=30, location="Virtual"):
    start = datetime.strptime(datetime_str, "%Y-%m-%d %H:%M")
    end = start + timedelta(minutes=duration_minutes)
    event = {
        'summary': title,
        'location': location,
        'start': {'dateTime': start.isoformat(), 'timeZone': 'America/Toronto'},
        'end': {'dateTime': end.isoformat(), 'timeZone': 'America/Toronto'},
    }
    calendar_service.events().insert(calendarId='primary', body=event).execute()

def full_productivity_flow(user_email, meeting_time, file_paths):
    docs = load_documents(file_paths)
    vectorstore = build_vectorstore(docs)
    prompt = (
        "Generate a comprehensive onboarding report with a welcome message, key goals, "
        "training dates, expectations, and action items from these files."
    )
    summary = generate_report_with_context(vectorstore, prompt)
    export_to_docx(summary, "onboarding_summary.docx")
    export_to_pdf(summary, "onboarding_summary.pdf")
    send_email(user_email, "Welcome to the Team!", summary)
    schedule_meeting("Onboarding Sync", meeting_time)
    return summary

def summarize_emails(file_paths):
    docs = load_documents(file_paths)
    vectorstore = build_vectorstore(docs)
    return generate_report_with_context(vectorstore, "Summarize key points and action items.")

def send_onboarding_email(email, body):
    send_email(email, "Welcome to the Team!", body)
    return f"✅ Email sent to {email}"

def schedule_onboarding_meeting(datetime_str, topic="Onboarding Sync"):
    schedule_meeting(topic, datetime_str)
    return f"📅 Meeting scheduled for {datetime_str}"

def generate_report(file_paths):
    docs = load_documents(file_paths)
    vectorstore = build_vectorstore(docs)
    prompt = "Create a structured report from these documents with findings and recommendations."
    return generate_report_with_context(vectorstore, prompt)

def fill_out_form(file_paths):
    docs = load_documents(file_paths)
    vectorstore = build_vectorstore(docs)
    prompt = "Extract structured data like names, dates, tasks, and emails in key-value format."
    return generate_report_with_context(vectorstore, prompt)

def follow_up_email(file_paths):
    docs = load_documents(file_paths)
    vectorstore = build_vectorstore(docs)
    prompt = "Generate a follow-up thank-you email based on this document's context."
    return generate_report_with_context(vectorstore, prompt)

